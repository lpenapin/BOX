from utils import (resolve_actor_via_sa_rest, 
                    parse_box_datetime_to_utc,
                    ensure_unique_filename,
                    sanitize_and_truncate_filename,
                    moving_to_downloadCompleted
)
from db_helpers import log_download_to_db
from notifications import send_email_alert


def download_from_box(
    client,
    box_folder_ids,
    download_base_path,
    downloaded_files,
    failed_files,
    db_config,
    *,
    log_to_db: bool = True,
    impersonate_user_id=None  # kept for compatibility; NOT used (service account works)
):
    """
    Service-account downloader.
    - Claim logic UNCHANGED.
    - Actor email via REST (current version uploader_display_name or modified_by).
    - Writes Actor fields to same Excel & DB.
    """
    import logging, re, time, os
    from pathlib import Path
    from datetime import datetime, timezone
    import openpyxl
    from openpyxl import Workbook
    from boxsdk.exception import BoxAPIException

    def _append_log_to_excel(excel_path: Path, log: dict):
        try:
            headers = [
                "Claim Number", "File Name", "File Type", "Description", "Source Folder",
                "Destination Folder", "File Size (MB)", "Download Start", "Download End",
                "Total Time (Seconds)", "Download Status", "Download Error Reason",
                "Business Area", "Creation Date",
                # Actor columns (same sheet)
                "Actor Name", "Actor Login"
            ]
            safe = {k: ("" if v is None else v) for k, v in log.items()}
            row = [
                safe.get("claim_number",""), safe.get("file_name",""), safe.get("file_type",""),
                safe.get("description",""), safe.get("source_box_folder",""),
                safe.get("destination_folder",""), safe.get("file_size_mb",""),
                str(safe.get("download_start","")), str(safe.get("download_end","")),
                safe.get("total_time_seconds",""), safe.get("download_status",""),
                safe.get("download_error_reason",""), safe.get("business_area",""),
                str(safe.get("creation_date","")),
                # actor
                safe.get("actor_name",""), safe.get("actor_login",""), #safe.get("actor_source",""),
            ]
            _file_name = safe.get("file_name","")
            excel_path.parent.mkdir(parents=True, exist_ok=True)
            if not excel_path.exists():
                wb = Workbook(); ws = wb.active; ws.title = "File Metadata"; ws.append(headers)
            else:
                wb = openpyxl.load_workbook(excel_path); ws = wb.active
                # ensure headers exist (idempotent)
                if ws.max_row < 1:
                    ws.append(headers)
                else:
                    existing = [c.value for c in ws[1]]
                    for h in headers:
                        if h not in existing:
                            ws.cell(row=1, column=len(existing)+1, value=h)
                            existing.append(h)
            ws.append(row)
            wb.save(excel_path)
            logging.info(f"✅ File {_file_name}, added to metadata file {excel_path}")
        except Exception as e:
            logging.error(f"❌ MetaData insert failed for {_file_name}: {e}")


    base_path = Path(download_base_path)

    # helpful: who am I?
    try:
        me = client.user('me').get()
        logging.info(f"👤 SA identity: id={me.id} login={getattr(me,'login',None)}")
    except Exception as e:
        logging.info(f"Could not determine SA identity: {e}")

    for box_folder_id in box_folder_ids:
        folder = client.folder(box_folder_id).get(fields=['name'])
        folder_name = folder.name.replace(" ", "_")
        raw_folder_base = base_path / folder_name
        raw_folder_base.mkdir(parents=True, exist_ok=True)

        logging.info(f"📁 Processing folder: {folder_name}")
        items = client.folder(box_folder_id).get_items(limit=1000, fields=['type','id','name'])

        for item in items:
            if getattr(item, "type", None) != "file":
                continue

            file_name = ""
            start_time = datetime.now(timezone.utc)
            try:
                file = client.file(item.id).get(fields=[
                    'id','name','description','owned_by','size',
                    'created_at','created_by','modified_at','modified_by'#,'file_version'
                ])
                file_name = file.name
                logging.info(f"📄 Found file: {file_name}")

                description = file.description or ""

                # ✅ YOUR EXACT CLAIM LOGIC (unchanged)
                claim_numbers = set(
                    re.findall(r"(?<!\d)(3\d{7})(?!\d)", description) 
                    + re.findall(r"(?<!\d)(3\d{7})(?!\d)", file_name) # CAN-6054
                )
                

                # 🔎 Resolve actor via the same SA REST path that worked in your probe
                actor = resolve_actor_via_sa_rest(client, file.id)
                actor_name   = actor['actor_name']
                actor_login  = actor['actor_login']   # email when available
                actor_source = actor['actor_source']

#CAN - 6054
                if len(claim_numbers) > 1:
                    subject = "🚨 BOX ALERT! - One file with two or more claim numbers"
                    body = f"📅 Date: {file.created_at}\n"
                    body += f'👤 Adjuster: {actor_name}\n'
                    body += f'📁 File Name: {file.name}\n'
                    body += f'#️⃣ Claim Numbers: {claim_numbers}\n'
                    body += f'🏢 LOB: {folder_name}\n'
                    send_email_alert(body, subject)
#END

                if not claim_numbers:
                    log = {
                        'file_name': file.name,
                        'claim_number': "",
                        'file_type': Path(file.name).suffix.lstrip('.').lower(),
                        'creation_date': parse_box_datetime_to_utc(getattr(file,'created_at',None)),
                        'description': description,
                        'source_box_folder': folder.name,
                        'destination_folder': str(raw_folder_base),
                        'file_size_mb': round((getattr(file,'size',0) or 0)/(1024*1024), 2),
                        'download_start': start_time,
                        'download_end': datetime.now(timezone.utc),
                        'total_time_seconds': 0,
                        'download_status': "Skipped",
                        'download_error_reason': "Missing claim number",
                        'business_area': folder_name,
                        # actor
                        'actor_name': actor_name,
                        'actor_login': actor_login,
                        'actor_source': actor_source,
                    }
                    excel_path = raw_folder_base / "missing_claims_metadata.xlsx"
                    _append_log_to_excel(excel_path, log)
                    if log_to_db:
                        try: log_download_to_db(db_config, log)  # your function
                        except Exception: pass
                    continue

                # 🔽 Process ALL matched claims
                for claim_number in claim_numbers:
                    raw_folder = raw_folder_base / claim_number
                    raw_folder.mkdir(parents=True, exist_ok=True)

                    safe_name = sanitize_and_truncate_filename(raw_folder, file_name)
                    unique_file_name = ensure_unique_filename(raw_folder, safe_name)
                    file_path = raw_folder / unique_file_name

                    # Download with retry
                    for attempt in range(3):
                        try:
                            with file_path.open('wb') as output:
                                client.file(file.id).download_to(output)
                            break
                        except BoxAPIException as e:
                            if e.status == 503 and attempt < 2:
                                logging.warning(f"🚧 503 for {file_name}. Retrying ({attempt+1}/3)… req_id={getattr(e,'request_id','N/A')}")
                                time.sleep(5)
                            else:
                                raise

                    end_time = datetime.now(timezone.utc)
                    downloaded_files.append(str(file_path))
                    logging.info(f"⬇️ Downloaded: {file_path}")

                    log = {
                        'file_name': file.name,
                        'claim_number': claim_number,
                        'file_type': file_path.suffix.lstrip('.').lower(),
                        'creation_date': parse_box_datetime_to_utc(getattr(file,'created_at',None)),
                        'description': description,
                        'source_box_folder': folder.name,
                        'destination_folder': str(raw_folder),
                        'file_size_mb': round((getattr(file,'size',0) or 0)/(1024*1024), 2),
                        'download_start': start_time,
                        'download_end': end_time,
                        'total_time_seconds': int((end_time - start_time).total_seconds()),
                        'download_status': "Success",
                        'download_error_reason': "",
                        'business_area': folder_name,
                        # actor
                        'actor_name': actor_name,
                        'actor_login': actor_login,
                        'actor_source': actor_source,
                    }

                    if log_to_db:
                        try:
                            ok = log_download_to_db(db_config, log)
                            if not ok:
                                logging.warning(f"⚠️ DB logging skipped for: {log['file_name']}")
                        except Exception as e:
                            logging.error(f"❌ DB insert failed for {file_name}: {e}")

                    excel_path = raw_folder / f"Xerox_{claim_number}_Metadata.xlsx"
                    _append_log_to_excel(excel_path, log)

#CAN - 6091
                # Funtion to move the file to Download complete folder
                suffix = 0
                while True:
                    code, text = moving_to_downloadCompleted(file.id, box_folder_id)
                    if code == 200:
                        logging.info(f'File {file.id} moved successfully!')
                        #print('File moved successfully!')
                        break
                    elif code == 409:
                        #rename file
                        org_name = file.name
                        base, ext = os.path.splitext(org_name)
                            # Get names of other files in the folder to check for duplicates
                        items = client.folder(folder.id).get_items(limit=None, offset=0)
                        existing_names = [item.name for item in items if item.id != file.id]

                        suffix += 1
                        new_name = f"{base}_{suffix}{ext}"
                        while new_name in existing_names:
                                suffix += 1
                                new_name = f"{base}_{suffix}{ext}"
                        #suffix += 1
                        
                        # Rename the file on Box
                        updated_file = client.file(file.id).update_info(data = {'name': new_name})
                        logging.info(f'File {file.id} renamed {updated_file} successfully!')
                    else:
                        print(f'Error: {code} - {text}')
                        logging.error(f'Error on file {file.id}: {code} - {text}')
                        break
#END

            except Exception as e:
                end_time = datetime.now(timezone.utc)
                failed_files.append(file_name or str(item.id))
                logging.error(f"❌ Error downloading file {file_name or item.id}: {e}")

                # Safe fallback if start_time not set
                start_time = start_time if 'start_time' in locals() else end_time

                log = {
                    'file_name': file_name or getattr(item, "name", str(item.id)),
                    'claim_number': "",
                    'file_type': Path(file_name or "").suffix.lstrip('.').lower() if file_name else "",
                    'creation_date': None,
                    'description': "",
                    'source_box_folder': folder.name,
                    'destination_folder': str(raw_folder_base),
                    'file_size_mb': 0.0,
                    'download_start': start_time,
                    'download_end': end_time,
                    'total_time_seconds': int((end_time - start_time).total_seconds()),
                    'download_status': "Failure",
                    'download_error_reason': str(e),
                    'business_area': folder_name,
                    'actor_name': None,
                    'actor_login': None,
                    'actor_source': "error",
                }
                excel_path = raw_folder_base / "failed_files.xlsx"
                _append_log_to_excel(excel_path, log)
                if log_to_db:
                    try: log_download_to_db(db_config, log)
                    except Exception: pass

    logging.info("✅ All files processed (actor via current version REST).")