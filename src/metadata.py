from pathlib import Path
from openpyxl import Workbook, load_workbook
import time
from datetime import datetime
import re
import logging

from db_helpers import ensure_final_metadata_table,log_metadata_to_db
from utils_phase2 import(
    _find_claim_dir_for_path,
    _read_actor_from_csv,
    _read_actor_from_xlsx,
    _actor_from_claim_dir,
    _safe_getsize
)
FINAL_FILENAME = "Xerox_Final_Metadata.xlsx"
FINAL_META_HEADERS = [
    "timestamp_start", "timestamp_end", "Time Taken (sec)",
    "claim_number", "action", "status", "reason",
    "source_path", "source_ext", "source_bytes",
    "target_path", "target_ext", "target_bytes",
    "pages_out",
    "actor_name", "actor_login",
]

class FinalMeta:
    def __init__(self, claims_root: Path, db_config: dict):
        self.root = Path(claims_root).resolve()
        self.db_config = db_config
        ensure_final_metadata_table(self.db_config)


    def _ensure_book(self, claim_dir: Path) -> Path:
        out = claim_dir / FINAL_FILENAME
        if not out.exists():
            wb = Workbook()
            ws = wb.active
            ws.title = "log"
            ws.append(FINAL_META_HEADERS)
            wb.save(out)
        else:
            wb = load_workbook(out)
            ws = wb.active
            existing = [c.value for c in next(
                ws.iter_rows(min_row=1, max_row=1))]
            if [e or "" for e in existing] != FINAL_META_HEADERS:
                ws = wb.create_sheet(f"log_v{len(wb.sheetnames)}")
                ws.append(FINAL_META_HEADERS)
                wb.active = ws
                wb.save(out)
        return out


    def record(self, *, action: str, status: str, source_path: str = "", target_path: str = "",
            reason: str = "", pages_out: int | None = None, t_start: float | None = None,
            t_end: float | None = None, actor_source_meta: Path | None = None):
        now = time.time()
        t0, t1 = t_start or now, t_end or now
        ts0, ts1 = datetime.fromtimestamp(t0).strftime("%Y-%m-%d %H:%M:%S.%f")[:-3], \
            datetime.fromtimestamp(t1).strftime("%Y-%m-%d %H:%M:%S.%f")[:-3]
        secs = round(max(0, t1-t0), 3)

        claim_dir = _find_claim_dir_for_path(
            self.root, target_path or source_path or self.root)
        claim_num = claim_dir.name if re.fullmatch(
            r"\d{8,}", claim_dir.name or "") else ""

        if actor_source_meta:
            try:
                actor_name, actor_login = (_read_actor_from_xlsx(actor_source_meta)
                                            if actor_source_meta.suffix.lower() in {".xlsx", ".xlsm"}
                                            else _read_actor_from_csv(actor_source_meta))
            except:
                actor_name = actor_login = None
        else:
            actor_name, actor_login = _actor_from_claim_dir(claim_dir)

        row_list = [
            ts0, ts1, secs, claim_num, action, status, reason,
            source_path, Path(source_path).suffix.lower(
            ) if source_path else "", _safe_getsize(source_path),
            target_path, Path(target_path).suffix.lower(
            ) if target_path else "", _safe_getsize(target_path),
            pages_out or "", actor_name or "", actor_login or ""
        ]

        # Excel write
        out = self._ensure_book(claim_dir)
        wb = load_workbook(out)
        ws = wb.active
        ws.append(row_list)
        wb.save(out)

        # DB write
        try:
            row_dict = {
                "TimestampStart": row_list[0], "TimestampEnd": row_list[1], "TimeTakenSec": row_list[2],
                "ClaimNumber": row_list[3], "Action": row_list[4], "Status": row_list[5], "Reason": row_list[6],
                "SourcePath": row_list[7], "SourceExt": row_list[8], "SourceBytes": row_list[9],
                "TargetPath": row_list[10], "TargetExt": row_list[11], "TargetBytes": row_list[12],
                "PagesOut": int(row_list[13]) if str(row_list[13]).isdigit() else None,
                "ActorName": row_list[14], "ActorLogin": row_list[15]
            }
            if self.db_config:
                log_metadata_to_db(self.db_config, row_dict)
        except Exception as e:
            logging.error(f"❌ Failed to log row to DB: {e}")

